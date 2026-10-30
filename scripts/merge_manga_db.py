"""調査結果 g*.json を1つの作品DB（CSV + XLSX）にまとめる。

使い方: python3 scripts/merge_manga_db.py <g*.jsonのあるフォルダ> docs/anime/manga_research
"""
import csv, json, sys, re
from pathlib import Path

src = Path(sys.argv[1])
out = Path(sys.argv[2])
out.mkdir(parents=True, exist_ok=True)

COLS = [
    ("id", "ID"), ("title", "作品名"), ("title_kana", "よみ"), ("category", "分類"),
    ("series_group", "シリーズ"), ("relation", "位置づけ"), ("author_story", "原作"),
    ("author_art", "作画"), ("publisher", "出版社"), ("magazine", "掲載誌"),
    ("run_period", "連載期間"), ("volumes", "巻数"), ("genre_tags", "ジャンル"),
    ("related_works", "関連作品"), ("themes", "主なテーマ"), ("psychology", "心理描写"),
    ("dialogue_style", "セリフの特徴"), ("direction_style", "演出の特徴"),
    ("gag_style", "ギャグの特徴"), ("flatup_priority", "FLATUP優先度"),
    ("homage_mode", "演出モード"), ("flatup_use", "FLATUP活用案"),
    ("copyright_caution", "著作権上の注意"), ("net_reputation", "ネットの評判"),
    ("sources", "情報源"), ("checked_date", "確認日"), ("status", "確認状態"),
]

def norm(t):
    return re.sub(r"[\s・!！?？〜~\-－]", "", t).lower()

rows, seen, dups = [], set(), []
for f in sorted(src.glob("g*.json")):
    for r in json.loads(f.read_text(encoding="utf-8")):
        key = norm(r.get("title", ""))
        if key in seen:
            dups.append(r.get("title"))
            continue
        seen.add(key)
        r["sources"] = " | ".join(r.get("sources") or [])
        r["_group"] = f.stem
        rows.append(r)

order = {"ヤンキー・不良": 0, "両方": 1, "格闘技・武道": 2}
rows.sort(key=lambda r: (order.get(r.get("category"), 3), r.get("series_group", ""), r["_group"]))
for i, r in enumerate(rows, 1):
    r["id"] = f"W{i:03d}"

with open(out / "works_database.csv", "w", encoding="utf-8-sig", newline="") as fh:
    w = csv.writer(fh)
    w.writerow([j for _, j in COLS])
    for r in rows:
        w.writerow([str(r.get(k, "")) for k, _ in COLS])

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
wb = Workbook()
ws = wb.active
ws.title = "作品DB"
ws.append([j for _, j in COLS])
for r in rows:
    ws.append([str(r.get(k, "")) for k, _ in COLS])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F2937")
widths = {"作品名": 28, "FLATUP活用案": 50, "著作権上の注意": 40, "情報源": 50}
for i, (_, j) in enumerate(COLS, 1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(j, 22)
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "C2"
ws.auto_filter.ref = ws.dimensions

s = wb.create_sheet("集計")
from collections import Counter
s.append(["区分", "値", "件数"])
for name, key in (("分類", "category"), ("確認状態", "status")):
    for v, n in Counter(r.get(key, "") for r in rows).most_common():
        s.append([name, v, n])
s.append(["合計", "", len(rows)])
wb.save(out / "works_database.xlsx")

stats = Counter(r.get("status") for r in rows)
print(json.dumps({"total": len(rows), "status": stats, "category": Counter(r.get("category") for r in rows),
                  "duplicates_dropped": dups, "with_sources": sum(1 for r in rows if r["sources"])},
                 ensure_ascii=False, indent=1))
